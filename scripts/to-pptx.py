#!/usr/bin/env python3
"""Pack the rendered slide PNGs into a .pptx, one full-bleed image per slide.

For venues that insist on PowerPoint. Text is not editable — the slides are
pictures — but they display exactly as rendered.

    pip install python-pptx
    python3 scripts/to-pptx.py 01
"""
import sys, pathlib
from pptx import Presentation
from pptx.util import Inches, Emu

root = pathlib.Path(__file__).resolve().parent.parent
target = sys.argv[1] if len(sys.argv) > 1 else '01'

matches = sorted((root / 'out').glob(f'{target}-*'))
if not matches:
    sys.exit(f'no rendered output for {target!r} — run scripts/build.sh {target} first')
deck = matches[0]
pngs = sorted((deck / 'png').glob('*.png'))
if not pngs:
    sys.exit(f'no PNGs in {deck / "png"} — run scripts/build.sh {target} first')

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9, matches the frame
blank = prs.slide_layouts[6]

for png in pngs:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), Emu(0), Emu(0),
                             width=prs.slide_width, height=prs.slide_height)

out = deck / f'praktychna-{target}.pptx'
prs.save(out)
print(f'{len(pngs)} slide(s) -> {out.relative_to(root)}')
